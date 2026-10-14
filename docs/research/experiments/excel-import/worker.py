"""엑셀 가져오기 메모리·시간 측정 작업자. 한 프로세스에서 한 번 실행하고 JSON 한 줄을 출력한다.

사용: python worker.py <mode> <파일> [key=value ...]
mode
  baseline      : import만 (기준 RSS)
  precheck      : precheck.py 사전 검사
  open_ro       : load_workbook(read_only=True, data_only=True, keep_links=False) 열기만
  scan_ro       : 위로 열고 첫 시트를 iter_rows(values_only)로 끝까지 훑기(행·열·셀 상한 검사 포함)
  open_normal   : read_only=False(일반 모드)로 열기만
  scan_formula  : data_only=False로 훑으며 수식 셀(data_type 'f') 개수 세기(수식 감지 비용)
  e2e_months    : 월 = 열 -> 월 = 행 전치 + 셀 검증 end-to-end (인수 item_cap=100|0, year=2026)
  e2e_table     : 월 = 행 표를 그대로 읽어 행마다 검증(전치 없음)
환경변수 OPENPYXL_DEFUSEDXML=False 로 defusedxml 자동 사용을 끌 수 있다(openpyxl 기능).
최대 RSS는 resource.getrusage(RUSAGE_SELF).ru_maxrss(KiB). 기준 RSS는 import 직후 값이다.
"""

import json
import os
import resource
import sys
import time
import tracemalloc
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

T0 = time.perf_counter()


def rss_mb():
    return resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024.0


import openpyxl  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.xml import DEFUSEDXML, LXML  # noqa: E402

import headers  # noqa: E402
import precheck as pc  # noqa: E402

BASE_RSS = rss_mb()
T_IMPORT = time.perf_counter() - T0

MAX_ROWS, MAX_COLS, MAX_CELLS = 100_000, 100, 2_000_000
ERROR_CAP = 200  # 05 §1 잘림 표시: details.fields는 200건에서 자르고 total은 전체 건수


def open_ro(path):
    return load_workbook(path, read_only=True, data_only=True, keep_links=False)


class Limit(Exception):
    pass


def a1(row, col):
    """1부터 시작하는 행·열 -> 셀 주소(A1)."""
    name = ""
    while col:
        col, r = divmod(col - 1, 26)
        name = chr(65 + r) + name
    return f"{name}{row}"


def to_number(v):
    """셀 값 -> (Decimal | None, 오류 코드 | None, 경고 코드 | None). 엑셀 표시 정밀도(15자리)로 정규화한다."""
    if v is None:
        return None, None, None
    if isinstance(v, bool):
        return None, "BAD_FORMAT", None
    if isinstance(v, (int, float)):
        d = Decimal(f"{v:.15g}") if isinstance(v, float) else Decimal(v)
        if not d.is_finite():
            return None, "BAD_FORMAT", None
        return d, None, None
    if isinstance(v, str):
        s = v.strip()
        if s in ("", "-", "–", "—"):
            return None, None, "BLANK_TOKEN"
        try:
            d = Decimal(s)
        except InvalidOperation:
            return None, "BAD_FORMAT", None
        if not d.is_finite() or len(d.as_tuple().digits) > 38:
            return None, "BAD_FORMAT", None
        return d, None, "NUMBER_AS_TEXT"
    return None, "BAD_FORMAT", None  # datetime 등


RULE_MIN, RULE_MAX, RULE_DECIMALS = Decimal(0), Decimal(10**7), 2


def check_rules(d):
    if d < RULE_MIN or d > RULE_MAX:
        return "OUT_OF_RANGE"
    exp = d.normalize().as_tuple().exponent
    if isinstance(exp, int) and -exp > RULE_DECIMALS:
        return "OUT_OF_RANGE"
    return None


def run_scan(wb, out):
    ws = wb[wb.sheetnames[0]]
    rows = cells = 0
    types = Counter()
    stop = "끝까지"
    for row in ws.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS + 1, values_only=True):
        rows += 1
        if rows > MAX_ROWS:
            stop = "행 상한 초과"
            break
        if len(row) > MAX_COLS and row[MAX_COLS] is not None:
            stop = "열 상한 초과"
            break
        n = len(row)
        cells += n
        if cells > MAX_CELLS:
            stop = "셀 상한 초과"
            break
        for v in row:
            types[type(v).__name__] += 1
    out.update(rows=rows, cells=cells, stop=stop, types=dict(types))


def run_scan_formula(path, out):
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    ws = wb[wb.sheetnames[0]]
    f = n = 0
    for row in ws.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS + 1):
        for c in row:
            n += 1
            if c.data_type == "f":
                f += 1
    wb.close()
    out.update(cells_seen=n, formula_cells=f)


def run_e2e_months(path, opt, out):
    """월 = 열(행 = 항목) -> 월 = 행(항목 = 열). 검증 오류는 셀 주소와 코드로 모은다(상한 200건, 전체 건수 유지)."""
    item_cap = int(opt.get("item_cap", 100))
    year_hint = int(opt["year"]) if "year" in opt else None
    t = time.perf_counter()
    wb = open_ro(path)
    out["t_open"] = round(time.perf_counter() - t, 3)
    ws = wb[wb.sheetnames[0]]
    if year_hint is None:
        year_hint, src = headers.suggest_year(ws.title, [])
        out["year_source"] = src
    errors, err_total = [], 0
    warn = Counter()

    def add_error(cell, code):
        nonlocal err_total
        err_total += 1
        if len(errors) < ERROR_CAP:
            errors.append((cell, code))

    header_row = None
    period_cols = {}  # 열 번호(1부터) -> 기간 첫날
    label_col = None
    items = []  # (표시 이름, 행 번호)
    values = {}  # (item index) -> {period: Decimal}
    rows_seen = 0
    t = time.perf_counter()
    for r_idx, row in enumerate(ws.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS + 1, values_only=True), start=1):
        rows_seen = r_idx
        if r_idx > MAX_ROWS:
            raise Limit("행 상한 초과")
        if header_row is None:
            if r_idx > 30:
                raise Limit("머리글 행을 찾지 못함(상단 30행)")
            cand = {}
            for c_idx, v in enumerate(row, start=1):
                d, _ = headers.parse_period_header(v, year_hint=year_hint)
                if d is not None and not isinstance(v, (int, float)):
                    cand[c_idx] = d
            if len(cand) >= 2:
                header_row, period_cols = r_idx, cand
                label_col = min(period_cols) - 3 if min(period_cols) > 3 else 1
            continue
        label = row[label_col - 1] if len(row) >= label_col else None
        if label is None or (isinstance(label, str) and not label.strip()):
            warn["BLANK_ROW"] += 1
            continue
        name = str(label).strip()
        if name in ("소계", "합계", "계", "총계"):
            warn["TOTAL_ROW_SKIPPED"] += 1
            continue
        if item_cap and len(items) >= item_cap:
            warn["ITEM_OVER_CAP_IGNORED"] += 1
            continue
        idx = len(items)
        items.append((name, r_idx))
        col_vals = {}
        for c_idx, period in period_cols.items():
            v = row[c_idx - 1] if len(row) >= c_idx else None
            d, err, w = to_number(v)
            if w:
                warn[w] += 1
            if err:
                add_error(a1(r_idx, c_idx), err)
                continue
            if d is not None:
                code = check_rules(d)
                if code:
                    add_error(a1(r_idx, c_idx), code)
                    continue
            col_vals[period] = d
        values[idx] = col_vals
    out["t_read_convert"] = round(time.perf_counter() - t, 3)
    wb.close()
    if header_row is None:
        raise Limit("머리글 행을 찾지 못함")
    # 전치: 기간 오름차순 한 행씩, 항목 순서대로 열
    t = time.perf_counter()
    periods = sorted(set(period_cols.values()))
    dup = len(period_cols) - len(periods)
    out_rows = [tuple(values[i].get(p) for i in range(len(items))) for p in periods]
    out["t_transpose"] = round(time.perf_counter() - t, 3)
    out.update(
        header_row=header_row, rows_scanned=rows_seen, items=len(items), periods=len(periods), dup_periods=dup,
        out_rows=len(out_rows), out_cells=len(out_rows) * len(items), errors_total=err_total,
        errors_listed=len(errors), truncated=err_total > len(errors), warnings=dict(warn), year=year_hint,
    )


def run_e2e_table(path, opt, out):
    """이미 월 = 행인 표: 머리글 행 + 데이터 행(기준월 날짜, 글자 2열, 숫자 열). 행마다 검증하고 목표 행 목록을 만든다."""
    wb = open_ro(path)
    ws = wb[wb.sheetnames[0]]
    rows_out, errors, err_total = [], [], 0
    seen_periods = Counter()
    ncols = 0
    for r_idx, row in enumerate(ws.iter_rows(max_row=MAX_ROWS + 2, max_col=MAX_COLS + 1, values_only=True), start=1):
        if r_idx == 1:
            ncols = len(row)
            continue
        if r_idx > MAX_ROWS + 1:
            raise Limit("행 상한 초과")
        d0 = row[0]
        period, why = headers.parse_period_header(d0)
        if period is None:
            err_total += 1
            if len(errors) < ERROR_CAP:
                errors.append((a1(r_idx, 1), why))
            continue
        rec = [period, str(row[1]).strip() if row[1] is not None else None, str(row[2]).strip() if row[2] is not None else None]
        for c_idx in range(3, len(row)):
            d, err, _ = to_number(row[c_idx])
            if err or (d is not None and check_rules(d)):
                err_total += 1
                if len(errors) < ERROR_CAP:
                    errors.append((a1(r_idx, c_idx + 1), err or "OUT_OF_RANGE"))
                continue
            rec.append(d)
        seen_periods[(rec[1], period)] += 1
        rows_out.append(tuple(rec))
    wb.close()
    out.update(rows=len(rows_out), ncols=ncols, errors_total=err_total, errors_listed=len(errors),
               dup_keys=sum(1 for v in seen_periods.values() if v > 1))


def main():
    mode, path = sys.argv[1], sys.argv[2]
    opt = dict(a.split("=", 1) for a in sys.argv[3:])
    out = {
        "mode": mode, "file": os.path.basename(path), "size_mib": round(os.path.getsize(path) / 1048576, 2),
        "defusedxml": bool(DEFUSEDXML), "lxml": bool(LXML), "openpyxl": openpyxl.__version__,
        "python": sys.version.split()[0], "base_rss_mb": round(BASE_RSS, 1), "import_s": round(T_IMPORT, 3),
    }
    if opt.get("tracemalloc") == "1":
        tracemalloc.start()
    t = time.perf_counter()
    try:
        if mode == "baseline":
            pass
        elif mode == "precheck":
            r = pc.precheck(path, os.path.getsize(path), pc.Limits())
            out.update(ok=r.ok, codes=r.codes)
        elif mode == "open_ro":
            wb = open_ro(path)
            out["sheets"] = len(wb.sheetnames)
            wb.close()
        elif mode == "scan_ro":
            wb = open_ro(path)
            out["t_open"] = round(time.perf_counter() - t, 3)
            run_scan(wb, out)
            wb.close()
        elif mode == "open_normal":
            wb = load_workbook(path, read_only=False, data_only=True, keep_links=False)
            out["sheets"] = len(wb.sheetnames)
            ws = wb[wb.sheetnames[0]]
            out["dims"] = ws.dimensions
        elif mode == "scan_formula":
            run_scan_formula(path, out)
        elif mode == "e2e_months":
            if opt.get("precheck", "1") == "1":
                r = pc.precheck(path, os.path.getsize(path), pc.Limits())
                if not r.ok:
                    raise Limit("사전 검사 거부: " + ",".join(r.codes))
            run_e2e_months(path, opt, out)
        elif mode == "e2e_table":
            run_e2e_table(path, opt, out)
        else:
            raise SystemExit(f"unknown mode {mode}")
        out["status"] = "ok"
    except Limit as e:
        out["status"] = f"limit: {e}"
    except Exception as e:  # noqa: BLE001
        out["status"] = f"{type(e).__name__}: {str(e)[:120]}"
    out["total_s"] = round(time.perf_counter() - t, 3)
    out["max_rss_mb"] = round(rss_mb(), 1)
    out["delta_rss_mb"] = round(rss_mb() - BASE_RSS, 1)
    if opt.get("tracemalloc") == "1":
        cur, peak = tracemalloc.get_traced_memory()
        out["tracemalloc_peak_mb"] = round(peak / 1048576, 1)
    print(json.dumps(out, ensure_ascii=False))


if __name__ == "__main__":
    main()
