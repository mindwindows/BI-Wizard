"""한 프로세스(=uvicorn 워커 1개) 안에서 파싱 N개를 스레드로 동시에 돌릴 때: 총 시간, 최대 RSS, 하트비트 지연(이벤트 루프/헬스체크 대리)."""
import sys, os, time, threading, resource, json
sys.path.insert(0, "/work/site")
from openpyxl import load_workbook
path, n = sys.argv[1], int(sys.argv[2])
def scan():
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    ws = wb[wb.sheetnames[0]]
    rows = 0
    for row in ws.iter_rows(max_row=100001, max_col=101, values_only=True):
        rows += 1
    wb.close()
    return rows
lat = []
stop = False
def hb():
    while not stop:
        t = time.perf_counter(); time.sleep(0.01); lat.append(time.perf_counter() - t - 0.01)
h = threading.Thread(target=hb, daemon=True); h.start()
t0 = time.perf_counter()
ths = [threading.Thread(target=scan) for _ in range(n)]
[t.start() for t in ths]; [t.join() for t in ths]
wall = time.perf_counter() - t0
stop = True; h.join()
lat.sort()
print(json.dumps({"file": os.path.basename(path), "threads": n, "wall_s": round(wall, 2), "maxrss_mb": round(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024, 1),
                  "heartbeat_max_delay_ms": round(max(lat) * 1000, 1), "heartbeat_p99_ms": round(lat[int(len(lat) * .99)] * 1000, 1)}))
