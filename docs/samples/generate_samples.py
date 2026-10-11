#!/usr/bin/env python3
"""BI-Wizard 합성 엑셀 샘플 생성기 (Gate ③ 4라운드, U-2 대체).

실제 회사 자료가 아니라 난수(seed 고정)로 만든 합성 데이터다. 같은 openpyxl 버전에서는 같은 내용이 나오도록
시드와 파일 속성의 시각을 고정한다.

필요: openpyxl (저장소 requirements에는 넣지 않았다. 이 스크립트 전용).
    pip install openpyxl==3.1.5
실행:
    python docs/samples/generate_samples.py            # docs/samples/ 에 xlsx 4개 생성
    python docs/samples/generate_samples.py <출력폴더>   # 다른 폴더에 생성
    python docs/samples/generate_samples.py --verify    # 생성 후 read_only로 다시 읽어 내용 검증

주의: openpyxl은 수식의 계산 결과(캐시 값)를 파일에 쓰지 않는다. 그러면 가져오기 시험에서 data_only=True로 읽을 때
수식 셀이 빈 값(None)이 되어 실제 Excel 파일과 달라진다. 그래서 이 스크립트는 수식 결과를 파이썬으로 계산해
시트 XML의 <v>에 직접 넣는다(Excel이 계산한 값이 아니다).
"""
import datetime
import io
import os
import random
import sys
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
FIXED = datetime.datetime(2026, 10, 11, 0, 0, 0)
SEED = 20261011
HERE = os.path.dirname(os.path.abspath(__file__))

TEAMS = ["영업1팀", "영업2팀", "고객지원팀"]
BRANCHES = ["서울", "부산", "대구", "광주", "대전"]


# ---------------------------------------------------------------- 공통 도구
def finalize(wb, path, cached):
    """wb를 저장하고, 수식 셀에 캐시 값을 넣고, zip 시각을 고정해 다시 쓴다.
    cached: {시트이름: {"E2": 값, ...}} (값이 str이면 문자열 결과, bool이면 불리언)."""
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.properties.creator = "BI-Wizard synthetic sample"
    wb.properties.lastModifiedBy = "BI-Wizard synthetic sample"
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    # 시트 이름 -> 파일 경로
    ET.register_namespace("", NS)
    wbx = ET.fromstring(zin.read("xl/workbook.xml"))
    rels = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
    rid2target = {r.get("Id"): r.get("Target") for r in rels}
    RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    name2file = {}
    for s in wbx.find(f"{{{NS}}}sheets"):
        t = rid2target[s.get(RNS)]
        name2file[s.get("name")] = t.lstrip("/") if t.startswith("/") else "xl/" + t
    file2name = {v: k for k, v in name2file.items()}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in sorted(zin.infolist(), key=lambda i: i.filename):
            data = zin.read(info.filename)
            sheet = file2name.get(info.filename)
            if sheet and cached.get(sheet):
                root = ET.fromstring(data)
                vals = cached[sheet]
                done = 0
                for c in root.iter(f"{{{NS}}}c"):
                    f = c.find(f"{{{NS}}}f")
                    if f is None:
                        continue
                    ref = c.get("r")
                    if ref not in vals:
                        raise SystemExit(f"캐시 값 없음: {sheet}!{ref}")
                    v = c.find(f"{{{NS}}}v")
                    if v is None:
                        v = ET.SubElement(c, f"{{{NS}}}v")
                    val = vals[ref]
                    if isinstance(val, bool):
                        c.set("t", "b")
                        v.text = "1" if val else "0"
                    elif isinstance(val, str):
                        c.set("t", "str")
                        v.text = val
                    else:
                        if "t" in c.attrib:
                            del c.attrib["t"]
                        v.text = repr(val)
                    done += 1
                if done != len(vals):
                    raise SystemExit(f"{sheet}: 수식 셀 {done}개, 캐시 값 {len(vals)}개가 맞지 않음")
                data = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + ET.tostring(root, encoding="utf-8")
            zi = zipfile.ZipInfo(info.filename, (2026, 10, 11, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zout.writestr(zi, data)


def hdr(ws, row, values, fill="DDEBF7"):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------------------------------------------------------------- 1. 팀 월별 실적
def make_team_monthly(path):
    r = random.Random(SEED + 1)
    wb = Workbook()
    ws = wb.active
    ws.title = "월별실적"
    hdr(ws, 1, ["월", "팀", "상품매출", "서비스매출", "합계"])
    cached = {}
    row = 2
    for m in range(1, 13):
        for t in TEAMS:
            a = r.randint(80, 400) * 10000
            b = r.randint(10, 150) * 10000
            ws.cell(row=row, column=1, value=datetime.date(2025, m, 1)).number_format = "yyyy-mm-dd"
            ws.cell(row=row, column=2, value=t)
            ws.cell(row=row, column=3, value=a).number_format = "#,##0"
            ws.cell(row=row, column=4, value=b).number_format = "#,##0"
            ws.cell(row=row, column=5, value=f"=C{row}+D{row}").number_format = "#,##0"
            cached.setdefault("월별실적", {})[f"E{row}"] = a + b
            row += 1
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    for col in "CDE":
        ws.column_dimensions[col].width = 14
    ws.freeze_panes = "A2"
    finalize(wb, path, cached)


# ---------------------------------------------------------------- 2. 병합 셀·제목 행·빈 행이 있는 지저분한 표
def make_messy(path):
    r = random.Random(SEED + 2)
    wb = Workbook()
    ws = wb.active
    ws.title = "현황"
    ws["A1"] = "2026년 상반기 지점별 판매 현황 (단위: 천원)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A1:F1")
    ws["A2"] = None  # 빈 행
    ws["A3"] = "작성일: 2026-07-01 / 작성자: 합성 샘플"
    ws.merge_cells("A3:F3")
    # 두 줄짜리 머리글: 병합 셀
    ws["A4"] = "지점"
    ws.merge_cells("A4:A5")
    ws["B4"] = "1분기"
    ws.merge_cells("B4:C4")
    ws["D4"] = "2분기"
    ws.merge_cells("D4:E4")
    ws["F4"] = "비고"
    ws.merge_cells("F4:F5")
    for col, v in zip("BCDE", ["매출", "비용", "매출", "비용"]):
        ws[f"{col}5"] = v
    for ref in ["A4", "B4", "D4", "F4", "B5", "C5", "D5", "E5"]:
        ws[ref].font = Font(bold=True)
        ws[ref].alignment = Alignment(horizontal="center", vertical="center")
    row = 6
    cached = {}
    for i, b in enumerate(BRANCHES):
        ws.cell(row=row, column=1, value=b)
        for j in range(2, 6):
            ws.cell(row=row, column=j, value=r.randint(100, 900) * 10).number_format = "#,##0"
        ws.cell(row=row, column=6, value=r.choice(["", "신규 점포", "리모델링", "", ""]) or None)
        row += 1
        if i == 2:  # 데이터 중간의 빈 행
            row += 1
    first, last = 6, row - 1
    ws.cell(row=row, column=1, value="합계")
    for j, col in enumerate("BCDE", 2):
        ws.cell(row=row, column=j, value=f"=SUM({col}{first}:{col}{last})").number_format = "#,##0"
    total_row = row
    row += 2
    ws.cell(row=row, column=1, value="※ 일부 지점은 3월 중 휴점.")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    # 캐시 값
    vals = {}
    for j, col in enumerate("BCDE", 2):
        vals[f"{col}{total_row}"] = sum((ws.cell(row=x, column=j).value or 0) for x in range(first, last + 1))
    cached["현황"] = vals
    ws.column_dimensions["A"].width = 14
    finalize(wb, path, cached)


# ---------------------------------------------------------------- 3. 여러 시트(요약/데이터/설명)
def make_multi_sheet(path):
    r = random.Random(SEED + 3)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "요약"
    ws1 = wb.create_sheet("데이터")
    ws2 = wb.create_sheet("설명")
    hdr(ws1, 1, ["주문번호", "주문일", "지점", "품목", "수량", "단가", "금액"])
    items = ["노트북", "모니터", "키보드", "마우스", "도킹스테이션"]
    prices = {"노트북": 1200000, "모니터": 300000, "키보드": 80000, "마우스": 40000, "도킹스테이션": 220000}
    rows = []
    for i in range(1, 31):
        b = r.choice(BRANCHES)
        it = r.choice(items)
        q = r.randint(1, 10)
        d = datetime.date(2026, 1, 1) + datetime.timedelta(days=r.randint(0, 180))
        rows.append((f"A{i:04d}", d, b, it, q, prices[it], q * prices[it]))
    cached = {"데이터": {}, "요약": {}}
    for i, (no, d, b, it, q, p, amt) in enumerate(rows, 2):
        ws1.cell(row=i, column=1, value=no)
        ws1.cell(row=i, column=2, value=d).number_format = "yyyy-mm-dd"
        ws1.cell(row=i, column=3, value=b)
        ws1.cell(row=i, column=4, value=it)
        ws1.cell(row=i, column=5, value=q)
        ws1.cell(row=i, column=6, value=p).number_format = "#,##0"
        ws1.cell(row=i, column=7, value=f"=E{i}*F{i}").number_format = "#,##0"
        cached["데이터"][f"G{i}"] = amt
    n = len(rows) + 1
    hdr(ws0, 1, ["지점", "주문 건수", "금액 합계"])
    for k, b in enumerate(BRANCHES, 2):
        ws0.cell(row=k, column=1, value=b)
        ws0.cell(row=k, column=2, value=f'=COUNTIF(데이터!C2:C{n},A{k})')
        ws0.cell(row=k, column=3, value=f'=SUMIF(데이터!C2:C{n},A{k},데이터!G2:G{n})').number_format = "#,##0"
        cached["요약"][f"B{k}"] = sum(1 for x in rows if x[2] == b)
        cached["요약"][f"C{k}"] = sum(x[6] for x in rows if x[2] == b)
    ws2["A1"] = "이 파일은 BI-Wizard Gate ④ 가져오기 시험용 합성 샘플입니다."
    ws2["A2"] = "시트 구성: 요약(수식으로 데이터 시트 집계), 데이터(주문 30건), 설명(이 시트)."
    ws2["A3"] = "가져오기 마법사는 시트 목록을 보여 주고 사용자가 한 시트를 고르게 해야 합니다."
    ws2.column_dimensions["A"].width = 80
    finalize(wb, path, cached)


# ---------------------------------------------------------------- 4. 날짜·소수·참거짓·선택 목록 혼합
def make_mixed_types(path):
    r = random.Random(SEED + 4)
    wb = Workbook()
    ws = wb.active
    ws.title = "주문"
    hdr(ws, 1, ["주문일", "주문일시", "코드", "수량", "단가", "할인율", "배송완료", "상태", "비고", "금액"])
    states = ["접수", "처리중", "완료", "취소"]
    dv = DataValidation(type="list", formula1='"접수,처리중,완료,취소"', allow_blank=True)
    ws.add_data_validation(dv)
    cached = {"주문": {}}
    for i in range(2, 22):
        d = datetime.date(2026, 3, 1) + datetime.timedelta(days=r.randint(0, 120))
        dt = datetime.datetime.combine(d, datetime.time(r.randint(8, 20), r.randint(0, 59), 0))
        q = r.randint(1, 50)
        p = round(r.uniform(100, 99999), 2)
        disc = r.choice([0, 0.05, 0.1, 0.125, 0.2])
        done = r.random() < 0.6
        st = r.choice(states)
        note = r.choice([None, None, "긴급", "재고 확인 필요", "고객 요청으로 변경"])
        ws.cell(row=i, column=1, value=d).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=2, value=dt).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=i, column=3, value=f"{r.randint(0, 9999):05d}").number_format = "@"  # 앞자리 0이 있는 문자열 코드
        ws.cell(row=i, column=4, value=q)
        ws.cell(row=i, column=5, value=p).number_format = "#,##0.00"
        ws.cell(row=i, column=6, value=disc).number_format = "0.0%"
        ws.cell(row=i, column=7, value=done)
        c = ws.cell(row=i, column=8, value=st)
        dv.add(c)
        ws.cell(row=i, column=9, value=note)
        ws.cell(row=i, column=10, value=f"=D{i}*E{i}*(1-F{i})").number_format = "#,##0.00"
        cached["주문"][f"J{i}"] = q * p * (1 - disc)
    # 모서리 값 한 줄: 큰 수, 지수, 음수, 아주 작은 수, 빈 칸
    i = 22
    ws.cell(row=i, column=1, value=datetime.date(2026, 12, 31)).number_format = "yyyy-mm-dd"
    ws.cell(row=i, column=3, value="00012").number_format = "@"
    ws.cell(row=i, column=4, value=-3)
    ws.cell(row=i, column=5, value=1.23456789012345e15)
    ws.cell(row=i, column=6, value=0.00001).number_format = "0.00%"
    ws.cell(row=i, column=7, value=False)
    ws.cell(row=i, column=8, value="접수")
    ws.cell(row=i, column=10, value=f"=D{i}*E{i}*(1-F{i})")
    cached["주문"][f"J{i}"] = -3 * 1.23456789012345e15 * (1 - 0.00001)
    for col, w in zip("ABCDEFGHIJ", [12, 18, 8, 8, 14, 9, 10, 10, 22, 16]):
        ws.column_dimensions[col].width = w
    finalize(wb, path, cached)


FILES = {
    "synthetic-team-monthly.xlsx": make_team_monthly,
    "synthetic-messy-merged.xlsx": make_messy,
    "synthetic-multi-sheet.xlsx": make_multi_sheet,
    "synthetic-mixed-types.xlsx": make_mixed_types,
}


# ---------------------------------------------------------------- 검증
def verify(outdir):
    from openpyxl import load_workbook

    ok = True

    def check(cond, msg):
        nonlocal ok
        print(("  OK  " if cond else "  실패 ") + msg)
        ok = ok and cond

    def ro(name):
        return load_workbook(os.path.join(outdir, name), read_only=True, data_only=True, keep_links=False)

    wb = ro("synthetic-team-monthly.xlsx")
    rows = list(wb["월별실적"].iter_rows(values_only=True))
    check(wb.sheetnames == ["월별실적"] and len(rows) == 37, "team-monthly: 시트 1개, 헤더 + 36행")
    check(rows[0] == ("월", "팀", "상품매출", "서비스매출", "합계"), "team-monthly: 한글 머리글")
    check(all(isinstance(x[0], datetime.datetime) and x[0].day == 1 for x in rows[1:]), "team-monthly: 월 열은 모두 월 첫날 날짜")
    check(all(x[4] == x[2] + x[3] for x in rows[1:]), "team-monthly: 합계 수식 열의 캐시 값 = 상품+서비스")
    wb.close()

    wb = ro("synthetic-messy-merged.xlsx")
    rows = list(wb["현황"].iter_rows(values_only=True))
    check(rows[0][0].startswith("2026년 상반기") and rows[1][0] is None, "messy: 1행 제목, 2행 빈 행")
    check(rows[3][1] == "1분기" and rows[3][2] is None and rows[4][1] == "매출", "messy: 병합 머리글(병합 범위의 오른쪽 셀은 빈 값, 아래 줄 머리글)")
    blank = [i for i, x in enumerate(rows, 1) if all(v is None for v in x)]
    check(len(blank) >= 3, f"messy: 완전히 빈 행 {blank}")
    tot = [x for x in rows if x[0] == "합계"][0]
    check(isinstance(tot[1], (int, float)) and tot[1] > 0, "messy: 합계 수식 캐시 값 존재")
    wb.close()
    wbn = load_workbook(os.path.join(outdir, "synthetic-messy-merged.xlsx"))  # 병합 범위는 일반 모드에서만 보인다
    check(len(wbn["현황"].merged_cells.ranges) >= 7, f"messy: 병합 범위 {len(wbn['현황'].merged_cells.ranges)}개(일반 모드로 확인)")

    wb = ro("synthetic-multi-sheet.xlsx")
    check(wb.sheetnames == ["요약", "데이터", "설명"], "multi-sheet: 시트 이름 요약/데이터/설명")
    s = list(wb["요약"].iter_rows(values_only=True))
    d = list(wb["데이터"].iter_rows(values_only=True))
    check(sum(x[1] for x in s[1:]) == 30 and sum(x[2] for x in s[1:]) == sum(x[6] for x in d[1:]), "multi-sheet: 요약 수식 캐시 값 합 = 데이터 시트 합")
    wb.close()

    wb = ro("synthetic-mixed-types.xlsx")
    rows = list(wb["주문"].iter_rows(values_only=True))
    check(isinstance(rows[1][0], datetime.datetime) and isinstance(rows[1][1], datetime.datetime), "mixed: 날짜·일시는 datetime으로 읽힘")
    check(all(isinstance(x[2], str) for x in rows[1:]) and rows[-1][2] == "00012", "mixed: 코드 열은 문자열(앞자리 0 보존)")
    check(isinstance(rows[1][4], float) and isinstance(rows[1][3], int), "mixed: 단가 소수(float), 수량 정수(int)")
    check(all(isinstance(x[6], bool) for x in rows[1:]), "mixed: 배송완료는 bool")
    check(set(x[7] for x in rows[1:]) <= {"접수", "처리중", "완료", "취소"}, "mixed: 상태는 선택 목록 값")
    check(any(x[8] is None for x in rows[1:]) and any(x[8] for x in rows[1:]), "mixed: 비고에 빈 칸과 값이 섞임")
    check(rows[-1][4] == 1.23456789012345e15 and rows[-1][3] == -3, "mixed: 큰 수·음수")
    check(all(isinstance(x[9], float) for x in rows[1:]), "mixed: 금액 수식 캐시 값 존재")
    wb.close()
    wbn = load_workbook(os.path.join(outdir, "synthetic-mixed-types.xlsx"))
    check(len(wbn["주문"].data_validations.dataValidation) == 1, "mixed: 데이터 유효성(선택 목록) 1개(일반 모드로 확인)")
    for name in FILES:
        size = os.path.getsize(os.path.join(outdir, name))
        check(size < 100 * 1024, f"{name}: {size} 바이트 (< 100KB)")
    return ok


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    outdir = args[0] if args else HERE
    os.makedirs(outdir, exist_ok=True)
    for name, fn in FILES.items():
        fn(os.path.join(outdir, name))
        print("생성:", os.path.join(outdir, name))
    if "--verify" in argv:
        print("검증:")
        if not verify(outdir):
            sys.exit(1)


if __name__ == "__main__":
    main(sys.argv[1:])
